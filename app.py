import streamlit as st
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import re

# Setup Halaman
st.set_page_config(page_title="Pemformat Ringkasan Matkul", layout="centered")
st.title("📄 Pemformat Dokumen Ringkasan Matkul")
st.caption("Ubah teks ringkasan menjadi file Word rapi (Max 3 Halaman, TNR 12pt, Spasi 1.5, Justify)")

# Form Input Data Diri, Nama File, & Teks Ringkasan
with st.form("form_pemformat"):
    st.subheader("Data Mahasiswa & Mata Kuliah")
    col1, col2 = st.columns(2)
    with col1:
        nama = st.text_input("Nama Lengkap", value="Muhammad Dzulfikri Rivai")
        npm = st.text_input("NPM", value="24013010138")
    with col2:
        matkul = st.text_input("Nama Mata Kuliah", placeholder="Isi Mata Kuliah")
        judul_bab = st.text_input("Judul Bab / Topik", placeholder="Isi Judul Bab")
    
    st.subheader("Pengaturan Simpan File")
    # Input kustom nama file
    custom_filename = st.text_input(
        "Nama File Output (.docx)", 
        value=f"24013010138_Muhammad Dzulfikri Rivai_"
    )
    
    st.subheader("Teks Ringkasan")
    isi_ringkasan = st.text_area("Paste Teks Ringkasan di Sini:", height=300, placeholder="Paste teks hasil ringkasan dari AI di sini...")
    
    submitted = st.form_submit_button("🎨 Rapikan & Buat File DOCX")

def add_formatted_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            sub_parts = re.split(r'(\*.*?\*)', part)
            for sub_part in sub_parts:
                if sub_part.startswith('*') and sub_part.endswith('*'):
                    run = paragraph.add_run(sub_part[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(sub_part)

def parse_markdown_to_docx(doc, text):
    lines = text.split('\n')
    for line in lines:
        line_str = line.strip()
        if not line_str or line_str.startswith('---') or line_str.startswith('==='):
            continue

        if re.match(r'^\d+\.\s+[^\d]', line_str):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line_str)
            run.bold = True

        elif re.match(r'^\d+\.\d+\.?\s+', line_str):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line_str)
            run.bold = True

        elif re.match(r'^\d+\.\d+\.\d+\.?\s+', line_str):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line_str)
            run.bold = True

        elif line_str.startswith('* ') or line_str.startswith('- '):
            clean_bullet = re.sub(r'^[*\-]\s*', '', line_str)
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_text(p, clean_bullet)

        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, line_str)

def generate_docx(nama, npm, matkul, judul_bab, isi_ringkasan):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    p_nama = doc.add_paragraph()
    p_nama.paragraph_format.space_before = Pt(0)
    p_nama.paragraph_format.space_after = Pt(0)
    p_nama.add_run(f"Nama\t\t: {nama}")

    p_npm = doc.add_paragraph()
    p_npm.paragraph_format.space_before = Pt(0)
    p_npm.paragraph_format.space_after = Pt(0)
    p_npm.add_run(f"NPM\t\t: {npm}")

    p_matkul = doc.add_paragraph()
    p_matkul.paragraph_format.space_before = Pt(0)
    p_matkul.paragraph_format.space_after = Pt(8)
    p_matkul.add_run(f"Mata Kuliah\t: {matkul}")

    p_judul = doc.add_paragraph()
    p_judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_judul.paragraph_format.space_before = Pt(0)
    p_judul.paragraph_format.space_after = Pt(10)
    
    run_judul1 = p_judul.add_run("RINGKASAN MATA KULIAH\n")
    run_judul1.bold = True
    
    run_judul2 = p_judul.add_run(f"{judul_bab.upper()}")
    run_judul2.bold = True

    parse_markdown_to_docx(doc, isi_ringkasan)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

if submitted:
    if not isi_ringkasan.strip():
        st.warning("Silakan paste teks ringkasan terlebih dahulu!")
    else:
        try:
            docx_file = generate_docx(nama, npm, matkul, judul_bab, isi_ringkasan)
            
            # Memastikan akhiran file .docx
            final_filename = custom_filename.strip()
            if not final_filename.endswith(".docx"):
                final_filename += ".docx"

            st.success(f"Dokumen Word '{final_filename}' siap didownload!")
            st.download_button(
                label=f"📥 Download {final_filename}",
                data=docx_file,
                file_name=final_filename,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        except Exception as e:
            st.error(f"Terjadi kesalahan: {e}")
